# This module builds the Service State Page workbook — Appetite version,
# layered on top of EITHER base version (BP-2.0 or pre-2.0).
#
# Appetite tables are always sourced from BP-2.0 (ServicePage.py), never
# pre-2.0, per the user's instruction. _ServiceAppetiteMixin holds the
# BP-2.0-sourced methods this program needs (4.B.1.e.(1)/4.E/4.F/4.G/4.H and
# their shared helpers) plus a _sheetSpecs() override that's idempotent
# (safe to mix onto either base):
#
#   - S Table 4.B.1.e.(1) (BB, Barber/Beauty/Spa Professional Liability)
#     already exists identically in both versions — the insert is a no-op
#     on both bases, included only for completeness/documentation.
#   - S Table 4.E (RSS) / 4.F (PSS) / 4.G (PSPL) / 4.H (MPVS): don't exist
#     in pre-2.0 at all — pure additions.
#
# On a BP-2.0 base, all five are already present, so every insert below is
# skipped since the codes already exist — same mixin works unmodified on
# both bases.

import pandas as pd
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .ServicePage import Service as ServiceBP20
from .ServicePageCurrent import Service as ServicePre20

_THIN_BORDER = Border(left=Side(style='thin', color='C1C1C1'), right=Side(style='thin', color='C1C1C1'),
                       top=Side(style='thin', color='C1C1C1'), bottom=Side(style='thin', color='C1C1C1'))


class _ServiceAppetiteMixin:
    # S Table 4.B.1.e.(1). Barber, Beauty, or Spa Professional Liability —
    # already present in both versions; kept here only so the mixin's
    # _sheetSpecs insert list is self-documenting (the insert is a no-op).
    def buildBarberProfLiab(self):
        barberProfLiab = self.buildDataFrame("BP7_ProfLiabarbersBeauticians_Rate")
        barberProfLiab['Occurrence'] = barberProfLiab['LiabilityLimit'].apply(lambda x: "${0:,.0f}".format(x))
        barberProfLiab['Aggregate'] = barberProfLiab['AggregateLimit'].apply(lambda x: "${0:,.0f}".format(x))
        barberProfLiab['Occurrence / Aggregate'] = barberProfLiab['Occurrence'] + ' / ' + barberProfLiab['Aggregate']
        pivotedBarberProf = barberProfLiab.pivot(index=['LiabilityLimit', 'Occurrence / Aggregate'], columns='ProfessionType', values='BaseRate').reset_index(['LiabilityLimit', 'Occurrence / Aggregate']). \
                rename(columns={'Barber': 'Each Barber', 'Beautician': 'Each Beautician', 'Manicurist': 'Each Manicurist'}).sort_values(by=['LiabilityLimit'])
        del pivotedBarberProf['LiabilityLimit']
        return pivotedBarberProf

    # S Table 4.E. Repair Services Specialized Endorsement — copied verbatim
    # from ServicePage.Service (BP-2.0).
    def buildRepairSpecializedEndorsement(self):
        endorsementCharge = self.buildDataFrame("BP7_MiscellaneousSpecializedEndorsement_Charges")
        rows = endorsementCharge[endorsementCharge['SpecializedEndorsementName'] == 'Repair Services Specialized Endorsement']
        charge = float(rows['EndorsementCharge'].iloc[0])
        return pd.DataFrame({"Base premium for each Service premises": ["${0:,.2f}".format(charge)]})

    def _formatRepairSpecializedEndorsement(self, ws):
        ws.merge_cells('A3:C3')
        ws.merge_cells('A4:C4')

    # S Table 4.F. Pet Services Specialized Endorsement — copied verbatim
    # from ServicePage.Service (build/format + the shared Mobile
    # Equipment/Business Income helpers also used by MPVS below).
    def buildPetSpecializedEndorsement(self):
        pssRate = self.buildDataFrame("BP7_PetServicesSpecialized")
        rate = pssRate.query('Constant == "Y"')['PetServicesSpecializedRate'].iloc[0]
        return pd.DataFrame({"Base premium per policy": ["${0:,.2f}".format(float(rate))]})

    def _buildMobileEquipment(self, petServicesType):
        mobileEquip = self.buildDataFrame("BP7_PetMobileServicesPetEquipment")
        filtered = mobileEquip.query('PetServicesType == @petServicesType').sort_values(by='MobileEquipmentCoverageLimit')
        filtered = filtered.rename(columns={'MobileEquipmentCoverageLimit': 'Limits', 'MobileEquipmentCoverageRate': 'Rate'}). \
                filter(items=['Limits', 'Rate'])
        filtered['Limits'] = filtered['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        filtered['Rate'] = filtered['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return filtered

    def _buildPSBusinessIncome(self, firstType, addlType, firstLabel):
        biTable = self.buildDataFrame("BP7_PetMobileServicesBusinessIncome")
        firstRows = biTable.query('MobileBusinessIncomeType == @firstType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileBusinessIncomeCoverageRate': firstLabel}). \
                filter(items=['Limits (BI)', firstLabel])
        addlRows = biTable.query('MobileBusinessIncomeType == @addlType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileBusinessIncomeCoverageRate': 'Each Additional'}). \
                filter(items=['Limits (BI)', 'Each Additional'])
        merged = pd.merge(firstRows, addlRows, how='outer', on='Limits (BI)').sort_values(by='Limits (BI)')
        for col in ('Limits (BI)', firstLabel, 'Each Additional'):
            merged[col] = merged[col].apply(lambda x: "${0:,.0f}".format(x))
        return merged

    def buildPSMobileEquipment(self):
        return self._buildMobileEquipment("Pet Services")

    def buildPSBusinessIncomeVehicle(self):
        return self._buildPSBusinessIncome("1st Pet Service Customized Vehicle", "Each Addl Pet Service Customized Vehicle", "1st Vehicle")

    def buildPSBusinessIncomeWorker(self):
        return self._buildPSBusinessIncome("1st Pet Service Worker", "Each Addl Pet Service Worker", "1st Worker")

    def buildVetMobileEquipment(self):
        return self._buildMobileEquipment("Veterinarian")

    def buildVetBusinessIncomeVehicle(self):
        return self._buildPSBusinessIncome("1st Veterinarian Customized Vehicle", "Each Addl Veterinarian Customized Vehicle", "1st Vehicle")

    def buildVetBusinessIncomeWorker(self):
        return self._buildPSBusinessIncome("1st Veterinarian", "Each Addl Veterinarian", "1st Worker")

    def _appendLabeledBlocks(self, ws, boldFont, font, blocks, blank_before_first=False):
        row = ws.max_row + 1
        max_col = 1
        for i, (label, df) in enumerate(blocks):
            if i > 0 or blank_before_first:
                row += 1  # blank separator row, left untouched (no border)
            label_row = row
            header_row = label_row + 1
            n_cols = len(df.columns)
            label_cell = ws.cell(row=label_row, column=1, value=label)
            label_cell.font = boldFont
            for col, name in enumerate(df.columns, start=1):
                ws.cell(row=header_row, column=col, value=name)
            for r_off, (_, data_row) in enumerate(df.iterrows()):
                for col, val in enumerate(data_row, start=1):
                    ws.cell(row=header_row + 1 + r_off, column=col, value=val)
            for col in range(1, n_cols + 1):
                cell = ws.cell(row=header_row, column=col)
                cell.font = boldFont
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            for r_off in range(len(df)):
                for col in range(1, n_cols + 1):
                    cell = ws.cell(row=header_row + 1 + r_off, column=col)
                    cell.font = font
                    cell.border = _THIN_BORDER
                    cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            max_col = max(max_col, n_cols)
            row = header_row + len(df) + 1
        for col in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].bestFit = True

    def _formatPSSplzdEndo(self, ws, boldFont, font):
        ws.merge_cells('A3:C3')
        ws.merge_cells('A4:C4')
        blocks = [
            ("Mobile Equipment", self.buildPSMobileEquipment()),
            ("Business Income (BI) per Customized Vehicle", self.buildPSBusinessIncomeVehicle()),
            ("Business Income (BI) per Worker", self.buildPSBusinessIncomeWorker()),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks, blank_before_first=True)

    # S Table 4.G. Pet Services Professional Liability — copied verbatim
    # from ServicePage.Service.buildPetServicePL (BP-2.0).
    def buildPetServicePL(self):
        psProfLiab = self.buildDataFrame("BP7_PetServicesProfessionalLiability").copy()
        psProfLiab['Limits'] = psProfLiab['PerOccurrenceAggregateLimitCode'].str.split('/').str[0].astype('int64')
        psProfLiab = psProfLiab.sort_values(by='Limits').rename(columns={'PetServicesProfessionalLiabilityRate': 'Rate'})
        psProfLiab['Limits'] = psProfLiab['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        psProfLiab['Rate'] = psProfLiab['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return psProfLiab.filter(items=['Limits', 'Rate'])

    # S Table 4.H. Mobile Pet and Veterinarian Services Endorsement — copied
    # verbatim from ServicePage.Service._formatMPVS (BP-2.0); generateWorksheet
    # is called with an EMPTY dataframe for this table code.
    def _formatMPVS(self, ws, boldFont, font):
        for i, (heading, mobile, vehicle, worker) in enumerate((
                ("Pet Services", self.buildPSMobileEquipment(), self.buildPSBusinessIncomeVehicle(), self.buildPSBusinessIncomeWorker()),
                ("Veterinarian Services", self.buildVetMobileEquipment(), self.buildVetBusinessIncomeVehicle(), self.buildVetBusinessIncomeWorker()))):
            ws.cell(row=ws.max_row + (2 if i else 1), column=1, value=heading).font = boldFont
            self._appendLabeledBlocks(ws, boldFont, font, [
                ("Mobile Equipment", mobile),
                ("Business Income (BI) per Customized Vehicle", vehicle),
                ("Business Income (BI) per Worker", worker),
            ], blank_before_first=True)

    def _sheetSpecs(self, Service):
        sheetSpecs = super()._sheetSpecs(Service)

        codes = {spec[0] for spec in sheetSpecs}
        if 'BB' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'ERP'), len(sheetSpecs)) + 1
            sheetSpecs.insert(insertAt, ('BB', 'S Table 4.B.1.e.(1). Barber, Beauty, or Spa Professional Liability', self.buildBarberProfLiab, False, True, None, None))

        codes = {spec[0] for spec in sheetSpecs}
        if 'RSS' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'FR'), len(sheetSpecs)) + 1
            sheetSpecs.insert(insertAt, ('RSS', 'S Table 4.E. Repair Services Specialized Endorsement', self.buildRepairSpecializedEndorsement, False, True, None, self._formatRepairSpecializedEndorsement))

        codes = {spec[0] for spec in sheetSpecs}
        if 'PSS' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'RSS'), len(sheetSpecs) - 1) + 1
            sheetSpecs.insert(insertAt, ('PSS', 'S Table 4.F. Pet Services Specialized Endorsement', self.buildPetSpecializedEndorsement, False, True, None,
                                          lambda ws: self._formatPSSplzdEndo(ws, Service.fontBold, Service.font)))

        codes = {spec[0] for spec in sheetSpecs}
        if 'PSPL' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'PSS'), len(sheetSpecs) - 1) + 1
            sheetSpecs.insert(insertAt, ('PSPL', 'S Table 4.G. Pet Services Professional Liability', self.buildPetServicePL, False, True, None, lambda ws: ws.insert_rows(3)))

        codes = {spec[0] for spec in sheetSpecs}
        if 'MPVS' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'PSPL'), len(sheetSpecs) - 1) + 1
            sheetSpecs.insert(insertAt, ('MPVS', 'S Table 4.H. Mobile Pet and Veterinarian Services Endorsement', lambda: pd.DataFrame(), False, False, None,
                                          lambda ws: self._formatMPVS(ws, Service.fontBold, Service.font)))

        return sheetSpecs


class Service(_ServiceAppetiteMixin, ServiceBP20):
    pass


class ServiceCurrent(_ServiceAppetiteMixin, ServicePre20):
    pass
