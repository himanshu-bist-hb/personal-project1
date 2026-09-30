# This module builds the Office State Page workbook — Appetite version,
# layered on top of EITHER base version (BP-2.0 or pre-2.0).
#
# Appetite tables are always sourced from BP-2.0 (OfficePage.py), never
# pre-2.0, per the user's instruction. _OfficeAppetiteMixin holds the
# BP-2.0-sourced methods this program needs (4.C.4.A, 4.F-4.P and their
# shared helpers) plus a _sheetSpecs() override that's idempotent (safe to
# mix onto either base): none of these table numbers exist in pre-2.0's own
# sheet list at all (it ends at 4.E, Franchise Upgrade Endorsement), so every
# insert below is a pure addition, never a replacement.
#
# On a BP-2.0 base, all twelve are already present, so every insert is
# skipped since the codes already exist — same mixin works unmodified on
# both bases.

import pandas as pd
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .OfficePage import Office as OfficeBP20
from .OfficePageCurrent import Office as OfficePre20

_THIN_BORDER = Border(left=Side(style='thin', color='C1C1C1'), right=Side(style='thin', color='C1C1C1'),
                       top=Side(style='thin', color='C1C1C1'), bottom=Side(style='thin', color='C1C1C1'))


class _OfficeAppetiteMixin:
    # --- Shared Veterinarian/Pet Services helpers (used by VSPL 4.C.4.A,
    # VS 4.L.3, PSS 4.O) — copied verbatim from OfficePage.Office (BP-2.0). ---
    def _buildVetSpecializedLiabByPet(self, petType, tab="BP7_VeterinarianSpecializedProfessional"):
        vetLiab = self.buildDataFrame(tab)
        filtered = vetLiab.query('PetType == @petType').copy()
        filtered['_occ'] = filtered['PerOccurrenceAggregateLimitCode'].str.split('/').str[0].astype('int64')
        filtered = filtered.sort_values(by='_occ')
        filtered['Limits'] = filtered['PerOccurrenceAggregateLimitCode'].apply(
            lambda x: "/".join("{0:,.0f}".format(int(p)) for p in x.split('/')))
        filtered['Rate'] = filtered['VeterinarianSpecializedLiabilityRate'].apply(lambda x: "${0:,.0f}".format(x))
        return filtered.filter(items=['Limits', 'Rate'])

    def buildVetSpecializedLiabHousehold(self):
        return self._buildVetSpecializedLiabByPet("HouseholdPet")

    def buildVetSpecializedLiabNonHousehold(self):
        return self._buildVetSpecializedLiabByPet("OtherThanHouseholdPet")

    def _buildMobileEquipment(self, petServicesType):
        mobileEquip = self.buildDataFrame("BP7_PetMobileServicesPetEquipment")
        filtered = mobileEquip.query('PetServicesType == @petServicesType').sort_values(by='MobileEquipmentCoverageLimit')
        filtered = filtered.rename(columns={'MobileEquipmentCoverageLimit': 'Limits', 'MobileEquipmentCoverageRate': 'Rate'}). \
                filter(items=['Limits', 'Rate'])
        filtered['Limits'] = filtered['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        filtered['Rate'] = filtered['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return filtered

    def buildVetSpecializedMobileEquip(self):
        return self._buildMobileEquipment("Veterinarian")

    def buildPSMobileEquipment(self):
        return self._buildMobileEquipment("Pet Services")

    def _buildVetSpecializedBusinessIncome(self, firstType, addlType, firstLabel):
        biTable = self.buildDataFrame("BP7_PetMobileServicesBusinessIncome")
        firstRows = biTable.query('MobileBusinessIncomeType == @firstType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileEquipmentCoverageRate': firstLabel,
                                'MobileBusinessIncomeCoverageRate': firstLabel}). \
                filter(items=['Limits (BI)', firstLabel])
        addlRows = biTable.query('MobileBusinessIncomeType == @addlType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileEquipmentCoverageRate': 'Each Additional',
                                'MobileBusinessIncomeCoverageRate': 'Each Additional'}). \
                filter(items=['Limits (BI)', 'Each Additional'])
        merged = pd.merge(firstRows, addlRows, how='outer', on='Limits (BI)').sort_values(by='Limits (BI)')
        for col in ('Limits (BI)', firstLabel, 'Each Additional'):
            merged[col] = merged[col].apply(lambda x: "${0:,.0f}".format(x))
        return merged

    def buildVetSpecializedBIVehicle(self):
        return self._buildVetSpecializedBusinessIncome("1st Veterinarian Customized Vehicle",
                                                       "Each Addl Veterinarian Customized Vehicle", "1st Vehicle")

    def buildVetSpecializedBIWorker(self):
        return self._buildVetSpecializedBusinessIncome("1st Veterinarian", "Each Addl Veterinarian", "1st Worker")

    def buildPSBusinessIncomeVehicle(self):
        return self._buildVetSpecializedBusinessIncome("1st Pet Service Customized Vehicle",
                                                       "Each Addl Pet Service Customized Vehicle", "1st Vehicle")

    def buildPSBusinessIncomeWorker(self):
        return self._buildVetSpecializedBusinessIncome("1st Pet Service Worker", "Each Addl Pet Service Worker", "1st Worker")

    # --- O Table 4.F-4.K: flat specialized-endorsement charges — copied
    # verbatim from OfficePage.Office (BP-2.0). ---
    def buildArchitectsEngineersEndorsement(self):
        return self._buildFlatSpecializedEndorsement('Architects and Engineers Specialized Endorsement',
                                                     header="Base premium for each Office premises")

    def _buildFlatSpecializedEndorsement(self, endorsementName, header="Base Premium for each Office Premises"):
        endorsementCharge = self.buildDataFrame("BP7_MiscellaneousSpecializedEndorsement_Charges")
        rows = endorsementCharge[endorsementCharge['SpecializedEndorsementName'] == endorsementName]
        charge = float(rows['EndorsementCharge'].iloc[0])
        return pd.DataFrame({header: ["${0:,.2f}".format(charge)]})

    def buildConsultantSpecializedEndorsement(self):
        return self._buildFlatSpecializedEndorsement('Consultants Specialized Endorsement',
                                                     header="Base premium for each risk premises")

    def buildProfessionalServicesEndorsement(self):
        return self._buildFlatSpecializedEndorsement('Professional Services Specialized Endorsement',
                                                     header="Base premium for each Office premises")

    def buildAccountantsSpecializedEndorsement(self):
        return self._buildFlatSpecializedEndorsement('Accountants Specialized Endorsement',
                                                     header="Base premium for each Office premises")

    def buildAttorneySpecializedEndorsement(self):
        return self._buildFlatSpecializedEndorsement('Attorneys Specialized Endorsement',
                                                     header="Base premium for each Office premises")

    def buildHealthCareSpecializedEndorsement(self):
        charges = self.buildDataFrame("BP7_HealthCareSpecialized_Charge")
        charges = charges[(charges['1to5EmployeesCharge'].fillna(0) > 0) & (charges['EachAddlEmployeeCharge'].fillna(0) > 0)]
        baseCharge = float(charges['1to5EmployeesCharge'].iloc[0])
        addlCharge = float(charges['EachAddlEmployeeCharge'].iloc[0])
        return pd.DataFrame({"Base premium for each Office premises": [
            "${0:,.2f}".format(baseCharge),
            "Plus ${0:,.0f} for each additional employee above 5,\npolicy wide for Employee Dishonesty coverage".format(addlCharge),
        ]})

    # --- O Table 4.L.3 (VS), 4.M.4 (VPL), 4.N (MPVS), 4.O (PSS), 4.P (PSPL)
    # build methods — copied verbatim from OfficePage.Office (BP-2.0). ---
    def buildVeterinarianSpecializedEndorsement(self):
        vsRate = self.buildDataFrame("BP7_VeterinarianSpecialized")
        rate = vsRate.query('Constant == "Y"')['VeterinarianSpecializedBaseChargeWithoutProfLiabRate'].iloc[0]
        return pd.DataFrame({"Base premium per policy": ["${0:,.2f}".format(float(rate))]})

    def buildVetProfLiabHousehold(self):
        return self._buildVetSpecializedLiabByPet("HouseholdPet", "BP7_VeterinarianProfessionalLiability")

    def buildVetProfLiabNonHousehold(self):
        return self._buildVetSpecializedLiabByPet("OtherThanHouseholdPet", "BP7_VeterinarianProfessionalLiability")

    def buildPetServicesMobileEquip(self):
        data = [("$15,000", "$49"), ("$25,000", "$85"), ("$50,000", "$166"), ("$100,000", "$220")]
        return pd.DataFrame(data, columns=["Limits", "Mobile Equipment"])

    def buildPetServicesCustomizedVehicle(self):
        data = [("$25,000", "$91", "$46"), ("$50,000", "$104", "$59"), ("$100,000", "$117", "$71")]
        return pd.DataFrame(data, columns=["Limits", "1st Vehicle", "Each Additional Vehicle"])

    def buildPetServicesBusinessIncome(self):
        data = [("$25,000", "$13", "$7"), ("$50,000", "$20", "$13"), ("$100,000", "$26", "$20")]
        return pd.DataFrame(data, columns=["Limits", "1st Worker", "Each Additional Worker"])

    def buildVetMobileEquip(self):
        data = [("$15,000", "$122"), ("$25,000", "$211"), ("$50,000", "$414"), ("$100,000", "$549")]
        return pd.DataFrame(data, columns=["Limits", "Mobile Equipment"])

    def buildVetCustomizedVehicle(self):
        data = [("$25,000", "$227", "$113"), ("$50,000", "$260", "$146"), ("$100,000", "$293", "$179")]
        return pd.DataFrame(data, columns=["Limits", "1st Vehicle", "Each Additional Vehicle"])

    def buildVetBusinessIncome(self):
        data = [("$25,000", "$32", "$17"), ("$50,000", "$49", "$32"), ("$100,000", "$65", "$49")]
        return pd.DataFrame(data, columns=["Limits", "1st Worker", "Each Additional Worker"])

    def buildPetServicesSpecializedEndorsement(self):
        pssRate = self.buildDataFrame("BP7_PetServicesSpecialized")
        rate = pssRate.query('Constant == "Y"')['PetServicesSpecializedRate'].iloc[0]
        return pd.DataFrame({"Base premium per policy": ["${0:,.2f}".format(float(rate))]})

    def buildPetServicesProfLiab(self):
        psProfLiab = self.buildDataFrame("BP7_PetServicesProfessionalLiability").copy()
        psProfLiab['Limits'] = psProfLiab['PerOccurrenceAggregateLimitCode'].str.split('/').str[0].astype('int64')
        psProfLiab = psProfLiab.sort_values(by='Limits').rename(columns={'PetServicesProfessionalLiabilityRate': 'Rate'})
        psProfLiab['Limits'] = psProfLiab['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        psProfLiab['Rate'] = psProfLiab['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return psProfLiab.filter(items=['Limits', 'Rate'])

    # --- Shared multi-block layout helper + per-sheet formatters — copied
    # verbatim from OfficePage.Office (BP-2.0). ---
    def _appendLabeledBlocks(self, ws, boldFont, font, blocks, blank_before_first=False):
        row = ws.max_row + 1
        max_col = 1
        for i, (label, df) in enumerate(blocks):
            if i > 0 or blank_before_first:
                row += 1  # blank separator row, left untouched (no border)
            label_row = row
            header_row = label_row + 1
            n_cols = len(df.columns)
            label_cell = ws.cell(row=label_row, column=1, value=label)
            label_cell.font = boldFont
            for col, name in enumerate(df.columns, start=1):
                ws.cell(row=header_row, column=col, value=name)
            for r_off, (_, data_row) in enumerate(df.iterrows()):
                for col, val in enumerate(data_row, start=1):
                    ws.cell(row=header_row + 1 + r_off, column=col, value=val)
            for col in range(1, n_cols + 1):
                cell = ws.cell(row=header_row, column=col)
                cell.font = boldFont
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            for r_off in range(len(df)):
                for col in range(1, n_cols + 1):
                    cell = ws.cell(row=header_row + 1 + r_off, column=col)
                    cell.font = font
                    cell.border = _THIN_BORDER
                    cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            max_col = max(max_col, n_cols)
            row = header_row + len(df) + 1  # next empty row, so the following block gets its blank separator
        for col in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].bestFit = True

    def _formatVSPL(self, ws, boldFont, font):
        self._appendLabeledBlocks(ws, boldFont, font, [
            ("Rate per Veterinarian - Household Pet", self.buildVetSpecializedLiabHousehold()),
            ("Rate per Veterinarian - Non Household Pet", self.buildVetSpecializedLiabNonHousehold()),
        ])
        noteRow = ws.max_row + 2
        ws.cell(row=noteRow, column=1, value="This coverage does not charge on the basis of per employee, but veterinarians only").font = font
        self._appendLabeledBlocks(ws, boldFont, font, [
            ("Mobile Equipment", self.buildVetSpecializedMobileEquip()),
            ("Business Income (BI) per Customized Vehicle", self.buildVetSpecializedBIVehicle()),
            ("Business Income (BI) per Worker", self.buildVetSpecializedBIWorker()),
        ], blank_before_first=True)
        ws.column_dimensions['A'].width = 160 / 7.0
        ws.column_dimensions['B'].width = 120 / 7.0
        ws.column_dimensions['C'].width = 120 / 7.0

    def _formatVS(self, ws, boldFont, font):
        ws.merge_cells('A3:C3')
        ws.merge_cells('A4:C4')
        blocks = [
            ("Mobile Equipment", self.buildVetSpecializedMobileEquip()),
            ("Business Income (BI) per Customized Vehicle", self.buildVetSpecializedBIVehicle()),
            ("Business Income (BI) per Worker", self.buildVetSpecializedBIWorker()),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks, blank_before_first=True)

    def _formatVPL(self, ws, boldFont, font):
        blocks = [
            ("Rate per Veterinarian - Household Pet", self.buildVetProfLiabHousehold()),
            ("Rate per Veterinarian - Non Household Pet", self.buildVetProfLiabNonHousehold()),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks)
        noteRow = ws.max_row + 2
        ws.cell(row=noteRow, column=1, value="This coverage does not charge on the basis of per employee, but veterinarians only").font = font

    def _formatMPVS(self, ws, boldFont, font):
        blocks = [
            ("Pet Services", self.buildPetServicesMobileEquip()),
            ("Pet Services per Customized Vehicle", self.buildPetServicesCustomizedVehicle()),
            ("Pet Services - Business Income", self.buildPetServicesBusinessIncome()),
            ("Veterinarian", self.buildVetMobileEquip()),
            ("Veterinarian Services per Customized Vehicle", self.buildVetCustomizedVehicle()),
            ("Veterinarian Services - Business Income", self.buildVetBusinessIncome()),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks)

    def _formatPSSplzdEndo(self, ws, boldFont, font):
        ws.merge_cells('A3:C3')
        ws.merge_cells('A4:C4')
        blocks = [
            ("Mobile Equipment", self.buildPSMobileEquipment()),
            ("Business Income (BI) per Customized Vehicle", self.buildPSBusinessIncomeVehicle()),
            ("Business Income (BI) per Worker", self.buildPSBusinessIncomeWorker()),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks, blank_before_first=True)

    # Idempotent insert-if-missing for each Appetite table, in rule-number
    # order. Every anchor code (OPTO/FR/AES/CS/PFSS/ACS/ATS/HCS/VS/VPL/MPVS/
    # PSS) is present on both bases (pre-2.0 either ships it natively, or an
    # earlier insert in this same call already added it), so `next(...)`
    # always finds a real match — the `len(sheetSpecs)` fallback only
    # guards against a future base-file reorder.
    def _sheetSpecs(self, Office):
        sheetSpecs = super()._sheetSpecs(Office)

        def insertAfter(anchorCode, spec):
            nonlocal sheetSpecs
            codes = {s[0] for s in sheetSpecs}
            if spec[0] in codes:
                return
            insertAt = next((i for i, s in enumerate(sheetSpecs) if s[0] == anchorCode), len(sheetSpecs) - 1) + 1
            sheetSpecs.insert(insertAt, spec)

        insertAfter('OPTO', ('VSPL', 'O Table 4.C.4.A. Veterinarian Specialized Endorsement With Professional Liability', lambda: pd.DataFrame(), False, False, None,
                              lambda ws: self._formatVSPL(ws, Office.fontBold, Office.font)))
        insertAfter('FR', ('AES', 'O Table 4.F. Architects and Engineers Specialized Endorsement', self.buildArchitectsEngineersEndorsement, False, True, None, None))
        insertAfter('AES', ('CS', 'O Table 4.G. Consultants Specialized Endorsement', self.buildConsultantSpecializedEndorsement, False, True, 'AES', None))
        insertAfter('CS', ('PFSS', 'O Table 4.H. Professional Services Specialized Endorsement', self.buildProfessionalServicesEndorsement, False, True, 'AES', None))
        insertAfter('PFSS', ('ACS', 'O Table 4.I. Accountants Specialized Endorsement', self.buildAccountantsSpecializedEndorsement, False, True, 'AES', None))
        insertAfter('ACS', ('ATS', 'O Table 4.J. Attorneys Specialized Endorsement', self.buildAttorneySpecializedEndorsement, False, True, 'AES', None))
        insertAfter('ATS', ('HCS', 'O Table 4.K. Health Care Specialized Endorsement', self.buildHealthCareSpecializedEndorsement, False, True, 'AES', None))
        insertAfter('HCS', ('VS', 'O Table 4.L.3. Veterinarian Specialized Endorsement', self.buildVeterinarianSpecializedEndorsement, False, True, 'PSS',
                             lambda ws: self._formatVS(ws, Office.fontBold, Office.font)))
        insertAfter('VS', ('VPL', 'O Table 4.M.4. Veterinarian Professional Liability', lambda: pd.DataFrame(), False, False, None,
                            lambda ws: self._formatVPL(ws, Office.fontBold, Office.font)))
        insertAfter('VPL', ('MPVS', 'O Table 4.N. Mobile Pet and Veterinarian Services Endorsement', lambda: pd.DataFrame(), False, False, None,
                             lambda ws: self._formatMPVS(ws, Office.fontBold, Office.font)))
        insertAfter('MPVS', ('PSS', 'O Table 4.O. Pet Services Specialized Endorsement', self.buildPetServicesSpecializedEndorsement, False, True, None,
                              lambda ws: self._formatPSSplzdEndo(ws, Office.fontBold, Office.font)))
        insertAfter('PSS', ('PSPL', 'O Table 4.P. Pet Services Professional Liability', self.buildPetServicesProfLiab, False, True, None, lambda ws: ws.insert_rows(3)))

        return sheetSpecs


class Office(_OfficeAppetiteMixin, OfficeBP20):
    pass


class OfficeCurrent(_OfficeAppetiteMixin, OfficePre20):
    pass
