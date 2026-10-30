# ==============================================================================
# CPHexValidate.py  —  Compare Territory-sheet factors in the SM / MM rate pages
#                      against the Hexbook (BGI / BGII / SCOL factors by grid).
# ------------------------------------------------------------------------------
# Function form of the desktop tool's CP_Hex_Validate.py: same comparisons and
# output workbook, minus the tkinter globals and the Excel sensitivity-label step
# (no other LOB applies one).
# ==============================================================================

import io
import os
from typing import Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

FACTOR_COLS = ["Territory", "BGI Factor", "BGII Factor", "SCOL Factor"]

# The Territory sheet holds three side-by-side blocks (Territory, Longitude,
# Latitude, BGI, BGII, SCOL) — these are the Territory + factor columns of each.
_BLOCKS = ([0, 3, 4, 5], [7, 10, 11, 12], [14, 17, 18, 19])


def _bytes(source) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if hasattr(source, "read"):
        return source.read()
    with open(str(source), "rb") as fh:
        return fh.read()


def _page_grid(page_bytes: bytes) -> pd.DataFrame:
    sheet = pd.read_excel(io.BytesIO(page_bytes), sheet_name="Territory", skiprows=2)
    blocks = []
    for cols in _BLOCKS:
        block = sheet.iloc[:, cols].copy()
        block.columns = FACTOR_COLS
        blocks.append(block)
    return pd.concat(blocks).dropna().reset_index(drop=True)


def _hex_grid(hex_bytes: bytes):
    book = pd.ExcelFile(io.BytesIO(hex_bytes))
    state = pd.read_excel(book, sheet_name="BGI").iloc[3, 1]
    grid = None
    for sheet, col in (("BGI", "BGI Factor"), ("BGII", "BGII Factor"), ("SCOL", "SCOL Factor")):
        part = pd.read_excel(book, sheet_name=sheet, skiprows=7, usecols=[0, 1], names=["Territory", col])
        grid = part if grid is None else grid.merge(part, on="Territory")
    return state, grid


def _diff(left: pd.DataFrame, right: pd.DataFrame, names) -> pd.DataFrame:
    a, b = left.set_index("Territory").align(right.set_index("Territory"))
    return a.compare(b, result_names=names).reset_index()


def _write_sheet(wb: Workbook, title: str, heading: str, diff: pd.DataFrame, left: str, right: str):
    ws = wb.create_sheet(title=title)
    ws["A1"] = heading
    ws["A2"] = f"Only shows differences between {left} & {right} Dataframes."
    ws["A3"] = "If there are no differences, the sheet is blank."
    for row in dataframe_to_rows(diff, False, True):
        ws.append(row)


def validate(hexbook, sm_pages=None, mm_pages=None, folder: Optional[str] = None) -> Dict:
    """
    Compare Hexbook factors with the SM and/or MM rate pages (at least one required).

    hexbook / sm_pages / mm_pages: bytes, file-like or path.
    Returns {"state", "path" (None when folder is None), "bytes", "counts"}, where
    counts maps each comparison sheet name to its number of differing territories.
    """
    if sm_pages is None and mm_pages is None:
        raise ValueError("Provide the SM and/or MM Excel rate page file.")

    state, hexgrid = _hex_grid(_bytes(hexbook))
    sm_grid = _page_grid(_bytes(sm_pages)) if sm_pages is not None else None
    mm_grid = _page_grid(_bytes(mm_pages)) if mm_pages is not None else None

    wb = Workbook()
    wb.active.title = "Hex Validation"
    wb["Hex Validation"]["A1"] = "This is the output of your comparison of rate page factors to our Hexbook factors."
    wb["Hex Validation"]["A2"] = ("The sheets outputted from this comparison tool will only highlight the rows in "
                                  "your Territory factor pages that deviate from Hexbook factors.")
    counts = {}

    def add(title, heading, left_name, left, right_name, right):
        diff = _diff(left, right, (left_name, right_name))
        counts[title] = int(len(diff))
        _write_sheet(wb, title, heading, diff, left_name, right_name)

    if sm_grid is not None:
        add("Hex to SM Comparison", "Hexbook to SM Page Comparison", "Hexbook", hexgrid, "SM Factors", sm_grid)
    if mm_grid is not None:
        add("Hex to MM Comparison", "Hexbook to MM Page Comparison", "Hexbook", hexgrid, "MM Factors", mm_grid)
    if sm_grid is not None and mm_grid is not None:
        add("SM to MM Comparison", "SM to MM Page Comparison", "SM Factors", sm_grid, "MM Factors", mm_grid)

    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()

    path = None
    if folder:
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, f"{state} Hex Factor Comparison.xlsx")
        try:
            with open(path, "wb") as fh:
                fh.write(data)
        except PermissionError:
            raise PermissionError(f"Close '{os.path.basename(path)}' in Excel and try again.")
    return {"state": str(state), "path": path, "bytes": data, "counts": counts}
