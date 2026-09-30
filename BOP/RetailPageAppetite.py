# This module builds the Retail State Page workbook — Appetite version,
# layered on top of EITHER base version (BP-2.0 or pre-2.0).
#
# Appetite tables are always sourced from BP-2.0 (RetailPage.py), never
# pre-2.0, per the user's instruction. _RetailAppetiteMixin holds the
# BP-2.0-sourced methods this program needs (4.C/4.E/4.F and their shared
# helpers) plus a _sheetSpecs() override that's idempotent (safe to mix onto
# either base):
#
#   - R Table 4.C: pre-2.0's own sheet list has "PLUS" (Retail PLUS
#     Endorsement) at this same table number — a DIFFERENT table than
#     BP-2.0's "RTS" (Retail Trade Specialized Endorsement). Per the user's
#     rule ("if pre2.0 has a table with some number and that table is also
#     in the appetite, provide the appetite table code instead"), PLUS is
#     removed and RTS is inserted in its place.
#   - R Table 4.E (PSS) / 4.F (PSPL): don't exist in pre-2.0 at all — pure
#     additions.
#
# On a BP-2.0 base, RTS/PSS/PSPL are already present (pre-2.0's PLUS never
# was), so the removal is a no-op and the insertions are skipped since the
# codes already exist — same mixin works unmodified on both bases.

import pandas as pd
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .RetailPage import Retail as RetailBP20
from .RetailPageCurrent import Retail as RetailPre20

_THIN_BORDER = Border(left=Side(style='thin', color='C1C1C1'), right=Side(style='thin', color='C1C1C1'),
                       top=Side(style='thin', color='C1C1C1'), bottom=Side(style='thin', color='C1C1C1'))


class _RetailAppetiteMixin:
    # R Table 4.C. Retail Trade Specialized Endorsement — copied verbatim
    # from RetailPage.Retail.buildRTSplzdEndo (BP-2.0).
    def buildRTSplzdEndo(self):
        endorsementCharge = self.buildDataFrame("BP7_MiscellaneousSpecializedEndorsement_Charges")
        return endorsementCharge[endorsementCharge['SpecializedEndorsementName'] == 'Retail Trade Specialized Endorsement'] \
            .filter(items=['EndorsementCharge']).rename(columns={'EndorsementCharge': 'Base premium for each Retail Premises'}).head(1)

    # R Table 4.E. Pet Services Specialized Endorsement — copied verbatim
    # from RetailPage.Retail (build/format + the three helper build methods
    # the format hook needs).
    def buildPSSplzdEndo(self):
        pssRate = self.buildDataFrame("BP7_PetServicesSpecialized")
        filteredPSSRate = pssRate.query('Constant == "Y"')
        rate = filteredPSSRate['PetServicesSpecializedRate'].iloc[0]
        return pd.DataFrame({"Base premium for each Retail Premises": ["${0:,.2f}".format(float(rate))]})

    def buildPSMobileEquipment(self):
        mobileEquip = self.buildDataFrame("BP7_PetMobileServicesPetEquipment")
        filteredMobileEquip = mobileEquip.query('PetServicesType == "Pet Services"').sort_values(by='MobileEquipmentCoverageLimit')
        filteredMobileEquip = filteredMobileEquip.rename(columns={'MobileEquipmentCoverageLimit': 'Limits', 'MobileEquipmentCoverageRate': 'Rate'}). \
                filter(items=['Limits', 'Rate'])
        filteredMobileEquip['Limits'] = filteredMobileEquip['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        filteredMobileEquip['Rate'] = filteredMobileEquip['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return filteredMobileEquip

    def _buildPSBusinessIncome(self, firstType, addlType, firstLabel):
        biTable = self.buildDataFrame("BP7_PetMobileServicesBusinessIncome")
        firstRows = biTable.query('MobileBusinessIncomeType == @firstType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileBusinessIncomeCoverageRate': firstLabel}). \
                filter(items=['Limits (BI)', firstLabel])
        addlRows = biTable.query('MobileBusinessIncomeType == @addlType'). \
                rename(columns={'MobileBusinessIncomeCoverageLimit': 'Limits (BI)', 'MobileBusinessIncomeCoverageRate': 'Each Additional'}). \
                filter(items=['Limits (BI)', 'Each Additional'])
        merged = pd.merge(firstRows, addlRows, how='outer', on='Limits (BI)').sort_values(by='Limits (BI)')
        merged['Limits (BI)'] = merged['Limits (BI)'].apply(lambda x: "${0:,.0f}".format(x))
        merged[firstLabel] = merged[firstLabel].apply(lambda x: "${0:,.0f}".format(x))
        merged['Each Additional'] = merged['Each Additional'].apply(lambda x: "${0:,.0f}".format(x))
        return merged

    def buildPSBusinessIncomeVehicle(self):
        return self._buildPSBusinessIncome("1st Pet Service Customized Vehicle", "Each Addl Pet Service Customized Vehicle", "1st Vehicle")

    def buildPSBusinessIncomeWorker(self):
        return self._buildPSBusinessIncome("1st Pet Service Worker", "Each Addl Pet Service Worker", "1st Worker")

    def _appendLabeledBlocks(self, ws, boldFont, font, blocks, blank_before_first=False):
        row = ws.max_row + 1
        max_col = 1
        for i, (label, df) in enumerate(blocks):
            if i > 0 or blank_before_first:
                row += 1  # blank separator row, left untouched (no border)
            label_row = row
            header_row = label_row + 1
            n_cols = len(df.columns)
            ws.cell(row=label_row, column=1, value=label)
            for col, name in enumerate(df.columns, start=1):
                ws.cell(row=header_row, column=col, value=name)
            for r_off, (_, data_row) in enumerate(df.iterrows()):
                for col, val in enumerate(data_row, start=1):
                    ws.cell(row=header_row + 1 + r_off, column=col, value=val)
            for col in range(1, n_cols + 1):
                cell = ws.cell(row=label_row, column=col)
                cell.font = boldFont
                cell.border = _THIN_BORDER
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if n_cols > 1:
                ws.merge_cells(start_row=label_row, start_column=1, end_row=label_row, end_column=n_cols)
            for col in range(1, n_cols + 1):
                header_cell = ws.cell(row=header_row, column=col)
                header_cell.font = boldFont
                header_cell.border = _THIN_BORDER
                header_cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            for r_off in range(len(df)):
                for col in range(1, n_cols + 1):
                    cell = ws.cell(row=header_row + 1 + r_off, column=col)
                    cell.font = font
                    cell.border = _THIN_BORDER
                    cell.alignment = Alignment(horizontal='center', vertical='bottom', wrap_text=True)
            max_col = max(max_col, n_cols)
            row = header_row + len(df) + 1
        for col in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].bestFit = True

    def _formatPSSplzdEndo(self, ws, boldFont, font, mobileDf, vehicleDf, workerDf):
        ws.merge_cells('A3:C3')
        ws.merge_cells('A4:C4')
        blocks = [
            ("Mobile Equipment", mobileDf),
            ("Business Income (BI) per Customized Vehicle", vehicleDf),
            ("Business Income (BI) per Worker", workerDf),
        ]
        self._appendLabeledBlocks(ws, boldFont, font, blocks, blank_before_first=True)

    # R Table 4.F. Pet Services Professional Liability — copied verbatim
    # from RetailPage.Retail.buildPSProfLiab (BP-2.0).
    def buildPSProfLiab(self):
        psProfLiab = self.buildDataFrame("BP7_PetServicesProfessionalLiability").copy()
        psProfLiab['Limits'] = psProfLiab['PerOccurrenceAggregateLimitCode'].str.split('/').str[0].astype('int64')
        psProfLiab = psProfLiab.sort_values(by='Limits').rename(columns={'PetServicesProfessionalLiabilityRate': 'Rate'})
        psProfLiab['Limits'] = psProfLiab['Limits'].apply(lambda x: "${0:,.0f}".format(x))
        psProfLiab['Rate'] = psProfLiab['Rate'].apply(lambda x: "${0:,.0f}".format(x))
        return psProfLiab.filter(items=['Limits', 'Rate'])

    def _sheetSpecs(self, Retail):
        sheetSpecs = super()._sheetSpecs(Retail)
        sheetSpecs = [spec for spec in sheetSpecs if spec[0] != 'PLUS']
        codes = {spec[0] for spec in sheetSpecs}

        if 'RTS' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'FR'), len(sheetSpecs))
            sheetSpecs.insert(insertAt, ('RTS', 'R Table 4.C. Retail Trade Specialized Endorsement', self.buildRTSplzdEndo, False, True, None, None))

        codes = {spec[0] for spec in sheetSpecs}
        if 'PSS' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'FR'), len(sheetSpecs)) + 1
            sheetSpecs.insert(insertAt, ('PSS', 'R Table 4.E. Pet Services Specialized Endorsement', self.buildPSSplzdEndo, False, True, None,
                                          lambda ws: self._formatPSSplzdEndo(ws, Retail.fontBold, Retail.font, self.buildPSMobileEquipment(),
                                                                              self.buildPSBusinessIncomeVehicle(), self.buildPSBusinessIncomeWorker())))

        codes = {spec[0] for spec in sheetSpecs}
        if 'PSPL' not in codes:
            insertAt = next((i for i, spec in enumerate(sheetSpecs) if spec[0] == 'PSS'), len(sheetSpecs) - 1) + 1
            sheetSpecs.insert(insertAt, ('PSPL', 'R Table 4.F. Pet Services Professional Liability', self.buildPSProfLiab, False, True, 'PED', None))

        return sheetSpecs


class Retail(_RetailAppetiteMixin, RetailBP20):
    pass


class RetailCurrent(_RetailAppetiteMixin, RetailPre20):
    pass
